import sys
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

if getattr(sys, "frozen", False):
    project_folder = Path(sys.executable).resolve().parent
else:
    project_folder = Path(__file__).resolve().parent

input_folder = project_folder / "input"
output_folder = project_folder / "output"

xlsx_files = list(input_folder.glob('*.xlsx'))

if not xlsx_files:
    print("В папке input нет Excel-файлов")
else:
    dfs = []

    required_columns = ["name", "amount", "city"]

    for file in xlsx_files:
        df = pd.read_excel(file)
        df.columns = df.columns.str.lower()

        missing_columns = [
            column for column in required_columns
            if column not in df.columns
        ]

        if missing_columns:
            print(
                f"Файл {file.name}: отсутствуют обязательные столбцы: "
                f"{', '.join(missing_columns)}"
            )
            exit()

        missing_amount_rows = df.index[df["amount"].isna()].tolist()

        if missing_amount_rows:
            rows = [row + 2 for row in missing_amount_rows]

            print(
                f"Файл {file.name}: найдены пропущенные значения в столбце amount."
            )
            print(f"Строки: {', '.join(map(str, rows))}")
            print("Обработка остановлена.")
            exit()

        dfs.append(df)

    orders = pd.concat(dfs, ignore_index=True)

    orders = orders.drop_duplicates()

    output_folder.mkdir(parents=True, exist_ok=True)

    number_of_orders = len(orders)
    total_amount = orders["amount"].sum()
    average_check = orders["amount"].mean().round(2)
    max_order = orders["amount"].max()
    max_order_name = orders.loc[orders["amount"].idxmax(), "name"]

    report = pd.DataFrame([
        {"Показатель": "Количество заказов", "Значение": number_of_orders},
        {"Показатель": "Общая сумма", "Значение": total_amount},
        {"Показатель": "Средний чек", "Значение": average_check},
        {"Показатель": "Самый крупный заказ", "Значение": max_order},
        {"Показатель": "Клиент с самым крупным заказом", "Значение": max_order_name},
    ])

    while True:
        try:
            large_order_limit = int(input("Введите минимальную сумму крупного заказа:"))

            if large_order_limit > 0:
                break
            else:
                print("Сумма должна быть больше 0")

        except ValueError:
            print("Нужно ввести целое число")

    city_groups = orders.groupby("city")

    city_report = city_groups["amount"].agg(["sum", "count", "mean"])
    city_report = city_report.reset_index()

    large_orders = orders[orders["amount"] > large_order_limit]

    sorted_orders = orders.sort_values("amount", ascending=False)

    if large_orders.empty:
        print("Крупных заказов не найдено")

    output_file = output_folder / "sales_report.xlsx"

    client_orders = orders[["name", "amount", "city"]].copy()
    client_orders.columns = ["Клиент", "Сумма", "Город"]

    client_large_orders = large_orders[["name", "amount", "city"]].copy()
    client_large_orders.columns = ["Клиент", "Сумма", "Город"]

    client_sorted_orders = sorted_orders[["name", "amount", "city"]].copy()
    client_sorted_orders.columns = ["Клиент", "Сумма", "Город"]

    client_report = report.copy()

    client_city_report = city_report.copy()
    client_city_report.columns = [
        "Город",
        "Общая сумма",
        "Количество заказов",
        "Средний чек"
    ]

    with pd.ExcelWriter(output_file) as writer:
        client_orders.to_excel(writer, sheet_name="Заказы", index=False)
        client_large_orders.to_excel(writer, sheet_name="Крупные заказы", index=False)
        client_sorted_orders.to_excel(writer, sheet_name="Заказы по сумме", index=False)
        client_report.to_excel(writer, sheet_name="Общий отчёт", index=False)
        client_city_report.to_excel(writer, sheet_name="Отчёт по городам", index=False)


    wb = load_workbook(output_file)

    # Стиль заливки
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for ws in wb.worksheets:
        # Выравнивание и границы
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
                cell.border = thin_border

        # Оформление заголовков
        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="1F1F1F"
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Закрепляем заголовок
        ws.freeze_panes = "A2"

        # Включаем фильтр
        ws.auto_filter.ref = ws.dimensions

        # Автоматическая ширина столбцов
        for column in ws.columns:
            max_length = 0

            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            width = max_length + 4
            ws.column_dimensions[column[0].column_letter].width = width

        # Форматы чисел для таблиц с заказами и отчёта по городам
        for cell in ws[1]:
            column_letter = cell.column_letter

            if cell.value in ("Сумма", "Общая сумма"):
                for row in range(2, ws.max_row + 1):
                    ws[f"{column_letter}{row}"].number_format = '#,##0  '

            elif cell.value == "Средний чек":
                for row in range(2, ws.max_row + 1):
                    ws[f"{column_letter}{row}"].number_format = '#,##0.00  '

            elif cell.value == "Количество заказов":
                for row in range(2, ws.max_row + 1):
                    ws[f"{column_letter}{row}"].number_format = '0  '

    # Отдельно форматируем значения в общем отчёте
    ws = wb["Общий отчёт"]

    for row in range(2, ws.max_row + 1):
        indicator = ws.cell(row=row, column=1).value
        value_cell = ws.cell(row=row, column=2)

        value_cell.font = Font(bold=True)

        if indicator == "Средний чек":
            value_cell.number_format = '#,##0.00  '

        elif indicator in (
                "Количество заказов",
                "Общая сумма",
                "Самый крупный заказ"
        ):
            value_cell.number_format = '#,##0  '

    # Сохраняем файл только после всех изменений
    wb.save(output_file)